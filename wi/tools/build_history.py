#!/usr/bin/env python3
"""Wisconsin general-election results 2016-2024 from the Wisconsin Elections Commission's official canvass reports
(elections.wi.gov; the Commission's site lets automated browsers read only its first results page, so 2016-2022 come from
the Internet Archive's copies of the same files). Writes, to OUT_DIR:
  wi_statewide_results_by_county.csv  race,year,office,county_fips,county,candidate,party,votes
  wi_house_results.csv / wi_senate_results.csv  year,race,plan,district,candidate,party,votes,unopposed   (house = State Assembly)
Every report has the same layout: office title, a header row ('Total Votes Cast' then party codes), a row of candidate
names, then one row per county (County by County) or ward (Ward by Ward, county named on its first ward), then 'Office Totals:'.
2020 President: Dane and Milwaukee counties use the recount results (certified Nov. 30, 2020).
Write-in candidates and SCATTERING are combined into one 'Write-ins' row (party W) per county or district.
Usage: build_history.py SRC_DIR OUT_DIR"""
import openpyxl, re, csv, sys, os, json, collections

SRC, OUT = sys.argv[1], sys.argv[2]
FIPS = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'wi_county_fips.json')))  # {county name: fips}
KEY = {re.sub(r'[^A-Z]', '', n.upper().replace('ST.', 'SAINT')): n for n in FIPS}
FILES = {2016: ['2016_cbc_president.xlsx', '2016_cbc_ussenate.xlsx', '2016_wbw_assembly.xlsx', '2016_wbw_senate.xlsx'],
         2018: ['2018_cbc_constitutional.xlsx', '2018_cbc_ussenate.xlsx', '2018_cbc_assembly.xlsx', '2018_cbc_senate.xlsx'],
         2020: ['2020_cbc_all.xlsx'],
         2022: ['2022_cbc_Governor_1.xlsx', '2022_cbc_Attorney_General_0.xlsx', '2022_cbc_Secretarry_of_State_0.xlsx', '2022_cbc_State_Treasurer_0.xlsx',
                '2022_cbc_US_Senate_0.xlsx', '2022_cbc_Representative_to_the_Assembly_0.xlsx', '2022_cbc_State_Senator_0.xlsx'],
         2024: ['2024_wbw_all.xlsx']}
# legislative maps in force: 2011 plan (Act 43) through 2020; the plan the Wisconsin Supreme Court adopted in April 2022
# (the Legislature's); the plan enacted in February 2024 (2023 Act 94) for 2024 and 2026
PLANS = {'house': {2016: 'a2011', 2018: 'a2011', 2020: 'a2011', 2022: 'a2022', 2024: 'a2024'},
         'senate': {2016: 's2011', 2018: 's2011', 2020: 's2011', 2022: 's2022', 2024: 's2024'}}
STATEWIDE = {'PRESIDENT OF THE UNITED STATES': 'President', 'UNITED STATES SENATOR': 'Senate', 'GOVERNOR/LIEUTENANT GOVERNOR': 'Governor',
             'GOVERNOR / LIEUTENANT GOVERNOR': 'Governor', 'GOVERNOR': 'Governor', 'ATTORNEY GENERAL': 'Attorney General',
             'SECRETARY OF STATE': 'Secretary of State', 'STATE TREASURER': 'Treasurer'}
PARTY = {'REP': 'R', 'DEM': 'D', 'LIB': 'L', 'WGR': 'G', 'GRN': 'G'}


def clean(s):
    s = str(s or '').split('\n')[0]            # tickets: keep the top of the ticket
    s = re.sub(r'\s+', ' ', s).strip()
    # one 2024 candidate is typed in capitals ('PRISCILLA A. PRADO'); names only, not county names or titles
    if re.fullmatch(r"[A-Z][A-Z .'-]+ [A-Z][A-Z.'-]+", s) and s.count(' ') <= 3 and not re.search(r'COUNTY|DISTRICT|TOTAL|SCATTERING|PRESIDENT|SENATOR|GOVERNOR|ATTORNEY|SECRETARY|TREASURER|REPRESENTATIVE', s):
        s = s.title()
    return s


def party_of(code, name):
    if not name or name.upper() == 'SCATTERING' or re.search(r'\((write[- ]?in|wr-in)\)', name, re.I):
        return 'W'
    return PARTY.get((code or '').strip().upper(), 'I')


def contests(path):
    """Yields (office title, [(name, party)], {county: [votes per candidate]}, [office totals]) for every sheet."""
    wb = openpyxl.load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows) if r and any(c == 'Total Votes Cast' for c in r)), None)
        if hi is None:
            continue
        tcol = rows[hi].index('Total Votes Cast')
        title = next(clean(r[0]) for r in reversed(rows[:hi]) if r and r[0] and clean(r[0]))
        hdr, names = rows[hi], rows[hi + 1]
        cols = [(j, clean(names[j]), party_of(hdr[j], re.sub(r'\s+', ' ', str(names[j])))) for j in range(tcol + 1, len(hdr)) if names[j] is not None and clean(names[j])]
        per, totals, county = collections.defaultdict(lambda: [0] * len(cols)), None, None
        for r in rows[hi + 2:]:
            if not r or not any(c is not None and c != '' for c in r):
                continue
            first = clean(r[0])
            if first.startswith('Office Totals'):
                totals = [int(r[j] or 0) for j, _, _ in cols]
                break
            if first:
                k = re.sub(r'[^A-Z]', '', first.upper().replace('ST.', 'SAINT'))
                if k not in KEY:
                    raise ValueError(f'{path} {title}: unknown county {first!r}')
                county = KEY[k]
            if county is None or clean(r[1] if len(r) > 1 else '').startswith('County Totals'):
                continue  # ward reports add a 'County Totals:' row after each county's wards
            for n, (j, _, _) in enumerate(cols):
                per[county][n] += int(r[j] or 0)
        yield title, [(n, p) for _, n, p in cols], per, totals




def recount_2020():
    """Dane and Milwaukee county totals after the Nov. 2020 presidential recount (the only counties recounted), from the
    Commission's final recount spreadsheet (2020 General Election Recount 11-30-20 final.xlsx): {county: {candidate: (original, recount)}}."""
    ws = openpyxl.load_workbook(os.path.join(SRC, '2020_recount_final.xlsx'), read_only=True)['Sheet1']
    rows = list(ws.iter_rows(values_only=True))
    names = [clean(str(x).split('/')[0]) for x in rows[1][5:10]]
    out = collections.defaultdict(lambda: {n: [0, 0] for n in names})
    for r in rows[2:]:
        c = clean(r[1]).upper().replace(' COUNTY', '')
        if not c:
            continue
        for i, n in enumerate(names):
            out[c.title()][n][0] += int(r[5 + i] or 0); out[c.title()][n][1] += int(r[11 + i] or 0)
    return out


RECOUNT = recount_2020()
sw_rows, dist_rows, checks = [], collections.defaultdict(list), []
for year, files in FILES.items():
    for f in files:
        for title, cands, per, totals in contests(os.path.join(SRC, f)):
            # county sums must equal the report's own office totals
            sums = [sum(v[n] for v in per.values()) for n in range(len(cands))]
            if totals is not None and sums != totals:
                checks.append(f'{year} {title}: county sums {sum(sums)} vs office totals {sum(totals)}')
            m = re.match(r'^(REPRESENTATIVE TO THE ASSEMBLY|STATE SENATOR) DISTRICT (\d+)$', title)
            if title in STATEWIDE:
                office = STATEWIDE[title]
                if year == 2020 and office == 'President':
                    for c, cand in RECOUNT.items():
                        for n, (nm, p) in enumerate(cands):
                            if nm in cand:
                                if per[c][n] != cand[nm][0]:
                                    checks.append(f'2020 recount: {c} {nm} canvass {per[c][n]} vs recount file original {cand[nm][0]}')
                                per[c][n] = cand[nm][1]
                race = f'{year} {office}'
                for c in sorted(FIPS):
                    if c not in per:
                        checks.append(f'{race}: no row for {c}')
                        continue
                    wi = 0
                    for n, (nm, p) in enumerate(cands):
                        if p == 'W':
                            wi += per[c][n]
                        else:
                            sw_rows.append([race, year, office, FIPS[c], c, nm, p, per[c][n]])
                    sw_rows.append([race, year, office, FIPS[c], c, 'Write-ins', 'W', wi])
            elif m:
                level = 'house' if m.group(1).startswith('REPRESENTATIVE') else 'senate'
                label = 'Assembly' if level == 'house' else 'Senate'
                d = int(m.group(2))
                real = [(nm, p, sums[n]) for n, (nm, p) in enumerate(cands) if p != 'W']
                unopp = int(len([1 for _, _, v in real if v > 0]) <= 1)
                for nm, p, v in real:
                    dist_rows[level].append([year, f'{year} {label}', PLANS[level][year], d, nm, p, v, unopp])
                wi = sum(sums[n] for n, (nm, p) in enumerate(cands) if p == 'W')
                if wi:
                    dist_rows[level].append([year, f'{year} {label}', PLANS[level][year], d, 'Write-ins', 'W', wi, unopp])
os.makedirs(OUT, exist_ok=True)
with open(os.path.join(OUT, 'wi_statewide_results_by_county.csv'), 'w', newline='') as fh:
    w = csv.writer(fh); w.writerow(['race', 'year', 'office', 'county_fips', 'county', 'candidate', 'party', 'votes']); w.writerows(sw_rows)
for level, rows in dist_rows.items():
    with open(os.path.join(OUT, f'wi_{level}_results.csv'), 'w', newline='') as fh:
        w = csv.writer(fh); w.writerow(['year', 'race', 'plan', 'district', 'candidate', 'party', 'votes', 'unopposed'])
        w.writerows(sorted(rows, key=lambda r: (r[0], r[3], r[5] == 'W', -r[6])))
print('statewide rows', len(sw_rows), {k: len(v) for k, v in dist_rows.items()})
for c in checks:
    print('check:', c)
